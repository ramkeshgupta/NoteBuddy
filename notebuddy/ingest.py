"""Turn uploaded files into clean text, and pick the most relevant passages for a question."""
from __future__ import annotations

import io
import math
import re
from collections import Counter
from pathlib import Path

SUPPORTED = ["pdf", "txt", "md", "docx"]

_STOP = set("""a an the and or but if of to in on at by for with from as is are was were be been it its this that
these those i you he she we they me my your our their what which who how why when where do does did can could
should would will about into than then so not no yes""".split())


def clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected.") from None
        text = "\n\n".join((p.extract_text() or "") for p in reader.pages)
    elif ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    elif ext in ("txt", "md", "markdown"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"'.{ext}' files aren't supported. Use PDF, DOCX, TXT or MD.")
    text = clean(text)
    if len(text) < 40:
        raise ValueError("No readable text found. Scanned PDFs (images) aren't supported yet.")
    return text


# ───────────────────────── retrieval ─────────────────────────
def chunk(text: str, size: int = 1000) -> list[str]:
    chunks, buf = [], ""
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        while len(para) > size:  # slice very long paragraphs on a sentence-ish boundary
            cut = para.rfind(". ", 0, size)
            cut = cut + 1 if cut > size // 2 else size
            if buf:
                chunks.append(buf)
                buf = ""
            chunks.append(para[:cut].strip())
            para = para[cut:].strip()
        if len(buf) + len(para) + 2 > size and buf:
            chunks.append(buf)
            buf = ""
        buf = f"{buf}\n\n{para}" if buf else para
    if buf:
        chunks.append(buf)
    return chunks


def _tokens(s: str) -> list[str]:
    return [t for t in re.findall(r"\w+", s.lower()) if t not in _STOP and len(t) > 1]


def _bm25(chunks: list[str], query: str, k: int) -> list[int]:
    q = set(_tokens(query))
    docs = [_tokens(c) for c in chunks]
    n = len(docs)
    if not q or n == 0:
        return list(range(min(k, n)))
    avg = (sum(len(d) for d in docs) / n) or 1
    df = Counter(t for d in docs for t in set(d))
    scored = []
    for i, d in enumerate(docs):
        tf = Counter(d)
        s = 0.0
        for t in q & tf.keys():
            idf = math.log(1 + (n - df[t] + 0.5) / (df[t] + 0.5))
            s += idf * tf[t] * 2.5 / (tf[t] + 1.5 * (0.25 + 0.75 * len(d) / avg))
        scored.append((s, i))
    scored.sort(reverse=True)
    top = [i for s, i in scored[:k] if s > 0]
    return top or list(range(min(k, n)))


def build_context(notes: list[dict], query: str, budget: int = 30000) -> str:
    """Return the notes as one text block. Small sets go in whole; big sets are narrowed to the best passages."""
    if not notes:
        return ""
    total = sum(len(n["text"]) for n in notes)
    if total <= budget:
        return "\n\n".join(f"=== {n['title']} ===\n{n['text']}" for n in notes)

    pool = []  # (note_index, chunk_index, text)
    for ni, n in enumerate(notes):
        for ci, c in enumerate(chunk(n["text"])):
            pool.append((ni, ci, c))
    order = _bm25([p[2] for p in pool], query, k=len(pool))
    picked, used = [], 0
    for idx in order:
        if used + len(pool[idx][2]) > budget:
            continue
        picked.append(pool[idx])
        used += len(pool[idx][2])
    picked.sort(key=lambda p: (p[0], p[1]))
    out, last = [], -1
    for ni, _, text in picked:
        if ni != last:
            out.append(f"=== {notes[ni]['title']} (relevant excerpts) ===")
            last = ni
        out.append(text)
    return "\n\n".join(out)
